"""Cara manual: baca Excel baris per baris dengan openpyxl, lalu hitung pakai dictionary."""
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

FILE = Path(__file__).with_name("KODE-WILAYAH-KEPMENDAGRI-2025.xlsx")

ws = load_workbook(FILE, read_only=True)["Sheet1"]
jumlah = Counter()
for baris in ws.iter_rows(min_row=2, values_only=True):
    tipe = baris[8]  # kolom "TIPE DESA(KELURAHAN, DESA, DESA ADAT)"
    jumlah[tipe] += 1

for tipe, n in jumlah.most_common():
    print(f"{tipe:<10} {n:>7,}".replace(",", "."))
